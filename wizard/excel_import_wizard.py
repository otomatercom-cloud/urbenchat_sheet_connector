import base64
import io
import logging

from odoo import models, fields, api, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class AdExcelImportWizard(models.TransientModel):
    """
    Import Ad IDs from a Meta Ads Excel export.

    Expected columns (any order, matched by header name):
        Campaign name  →  urbenchat.meta.campaign
        Ad set name    →  urbenchat.meta.adset
        Ad name        →  lead.source.campaign  (auto-created if missing)
        Ad ID          →  urbenchat.ad.mapping.ad_id

    Lead Source is auto-filled from the Source Campaign's parent Lead Source.
    """
    _name = 'urbenchat.excel.import.wizard'
    _description = 'Ad ID Excel Import Wizard'

    # ── Step 1: Upload ─────────────────────────────────────────────────────
    state = fields.Selection(
        [('upload', 'Upload'), ('preview', 'Preview'), ('done', 'Done')],
        default='upload',
    )
    excel_file    = fields.Binary(string='Excel / CSV File', attachment=False)
    file_name     = fields.Char(string='File Name')
    sheet_name    = fields.Char(
        string='Sheet Name',
        default='Formatted Report',
        help='Sheet tab name. Leave blank to use the first sheet.',
    )

    # Default Lead Source applied to all auto-created Source Campaigns
    default_leads_source_id = fields.Many2one(
        'leads.sources', string='Default Lead Source',
        required=True,
        help='All Source Campaigns created from the Ad name column will be '
             'placed under this Lead Source.',
    )

    # ── Step 2: Preview counts ──────────────────────────────────────────────
    preview_text     = fields.Text(string='Preview', readonly=True)
    rows_total       = fields.Integer(string='Total Rows', readonly=True)
    campaigns_new    = fields.Integer(string='New Campaigns', readonly=True)
    adsets_new       = fields.Integer(string='New Ad Sets', readonly=True)
    source_camps_new = fields.Integer(string='New Source Campaigns', readonly=True)
    ads_new          = fields.Integer(string='New Ad IDs', readonly=True)
    ads_skip         = fields.Integer(string='Duplicate Ad IDs (skip)', readonly=True)

    # ── Step 3: Result ──────────────────────────────────────────────────────
    result_text = fields.Text(string='Import Result', readonly=True)

    # ── Parse helpers ───────────────────────────────────────────────────────
    def _parse_file(self):
        if not self.excel_file:
            raise UserError(_('Please upload a file.'))
        raw = base64.b64decode(self.excel_file)
        fname = (self.file_name or '').lower()

        if fname.endswith('.csv'):
            import csv, io as _io
            text = raw.decode('utf-8-sig')
            reader = csv.DictReader(_io.StringIO(text))
            return list(reader)

        # Excel
        try:
            import openpyxl
        except ImportError:
            raise UserError(_(
                "openpyxl is required. Install it: pip install openpyxl --break-system-packages"
            ))
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        if self.sheet_name and self.sheet_name in wb.sheetnames:
            ws = wb[self.sheet_name]
        else:
            ws = wb.active

        rows_raw = list(ws.iter_rows(values_only=True))
        if not rows_raw:
            return []
        headers = [str(h).strip() if h else '' for h in rows_raw[0]]
        result  = []
        for row in rows_raw[1:]:
            if not any(row):
                continue
            result.append({headers[i]: (str(row[i]).strip() if row[i] is not None else '')
                           for i in range(len(headers))})
        return result

    def _get_col(self, row, *keys):
        """Try multiple possible column names."""
        for k in keys:
            v = row.get(k, '').strip()
            if v:
                return v
        return ''

    # ── Actions ─────────────────────────────────────────────────────────────
    def action_preview(self):
        self.ensure_one()
        if not self.default_leads_source_id:
            raise UserError(_('Please set a Default Lead Source before previewing.'))

        rows = self._parse_file()
        if not rows:
            raise UserError(_('File is empty or could not be parsed.'))

        AdMap    = self.env['urbenchat.ad.mapping']
        Camp     = self.env['urbenchat.meta.campaign']
        Adset    = self.env['urbenchat.meta.adset']
        SrcCamp  = self.env['lead.source.campaign']

        seen_camps   = set()
        seen_adsets  = set()
        seen_srccamps= set()
        seen_adids   = set()
        skip_adids   = set()

        existing_adids = set(AdMap.search([]).mapped('ad_id'))

        for row in rows:
            camp_name  = self._get_col(row, 'Campaign name', 'campaign_name', 'Campaign')
            adset_name = self._get_col(row, 'Ad set name', 'adset_name', 'Ad set')
            ad_name    = self._get_col(row, 'Ad name', 'ad_name', 'Ad Name')
            ad_id      = self._get_col(row, 'Ad ID', 'ad_id', 'Ad Id')

            if not ad_id:
                continue
            if ad_id in existing_adids or ad_id in seen_adids:
                skip_adids.add(ad_id)
                continue

            seen_adids.add(ad_id)
            if camp_name:  seen_camps.add(camp_name)
            if adset_name: seen_adsets.add(adset_name)
            if ad_name:    seen_srccamps.add(ad_name)

        # Count what's new
        existing_camps    = set(Camp.search([]).mapped('name'))
        existing_adsets   = set(Adset.search([]).mapped('name'))
        existing_srccamps = set(SrcCamp.search([]).mapped('name'))

        new_camps    = seen_camps    - existing_camps
        new_adsets   = seen_adsets   - existing_adsets
        new_srccamps = seen_srccamps - existing_srccamps

        self.rows_total       = len(rows)
        self.campaigns_new    = len(new_camps)
        self.adsets_new       = len(new_adsets)
        self.source_camps_new = len(new_srccamps)
        self.ads_new          = len(seen_adids)
        self.ads_skip         = len(skip_adids)

        lines = [
            f'Total rows in file:          {len(rows)}',
            f'Ad IDs to import:            {len(seen_adids)}',
            f'Duplicate Ad IDs (skip):     {len(skip_adids)}',
            '',
            f'New Meta Campaigns:          {len(new_camps)}',
            f'New Meta Ad Sets:            {len(new_adsets)}',
            f'New Source Campaigns (Odoo): {len(new_srccamps)}',
            '',
            '--- Sample (first 5 rows) ---',
        ]
        for row in rows[:5]:
            camp  = self._get_col(row, 'Campaign name', 'campaign_name', 'Campaign')
            adset = self._get_col(row, 'Ad set name', 'adset_name', 'Ad set')
            adnm  = self._get_col(row, 'Ad name', 'ad_name', 'Ad Name')
            adid  = self._get_col(row, 'Ad ID', 'ad_id', 'Ad Id')
            lines.append(f'  [{adid}]')
            lines.append(f'    Campaign:        {camp}')
            lines.append(f'    Ad Set:          {adset}')
            lines.append(f'    Source Campaign: {adnm}')

        self.preview_text = '\n'.join(lines)
        self.state = 'preview'
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        rows = self._parse_file()

        Camp     = self.env['urbenchat.meta.campaign']
        Adset    = self.env['urbenchat.meta.adset']
        SrcCamp  = self.env['lead.source.campaign']
        AdMap    = self.env['urbenchat.ad.mapping']
        src_lead = self.default_leads_source_id

        existing_adids = set(AdMap.search([]).mapped('ad_id'))

        # Caches
        camp_cache   = {c.name: c for c in Camp.search([])}
        adset_cache  = {a.name: a for a in Adset.search([])}
        srcamp_cache = {s.name: s for s in SrcCamp.search([])}

        created = skipped = errors = 0
        error_lines = []

        for row in rows:
            camp_name  = self._get_col(row, 'Campaign name', 'campaign_name', 'Campaign')
            adset_name = self._get_col(row, 'Ad set name', 'adset_name', 'Ad set')
            ad_name    = self._get_col(row, 'Ad name', 'ad_name', 'Ad Name')
            ad_id      = self._get_col(row, 'Ad ID', 'ad_id', 'Ad Id')

            if not ad_id:
                continue
            if ad_id in existing_adids:
                skipped += 1
                continue

            try:
                # 1. Get/create Meta Campaign
                if camp_name not in camp_cache:
                    camp_cache[camp_name] = Camp.create({'name': camp_name})
                campaign = camp_cache[camp_name]

                # 2. Get/create Meta Ad Set (scoped to campaign)
                adset_key = f'{campaign.id}::{adset_name}'
                if adset_key not in adset_cache:
                    existing = Adset.search([
                        ('name', '=', adset_name),
                        ('campaign_id', '=', campaign.id),
                    ], limit=1)
                    adset_cache[adset_key] = existing or Adset.create({
                        'name': adset_name,
                        'campaign_id': campaign.id,
                    })
                adset = adset_cache[adset_key]

                # 3. Get/create Source Campaign (Ad name → lead.source.campaign)
                if ad_name not in srcamp_cache:
                    existing_sc = SrcCamp.search([('name', '=', ad_name)], limit=1)
                    srcamp_cache[ad_name] = existing_sc or SrcCamp.create({
                        'name': ad_name,
                        'lead_source_id': src_lead.id,
                    })
                source_campaign = srcamp_cache[ad_name]

                # 4. Create Ad Mapping
                AdMap.create({
                    'adset_id':           adset.id,
                    'ad_id':              ad_id,
                    'source_campaign_id': source_campaign.id,
                })
                existing_adids.add(ad_id)
                created += 1

            except Exception as e:
                errors += 1
                error_lines.append(f'❌ Ad ID {ad_id}: {e}')
                _logger.error('Excel import error for ad_id %s: %s', ad_id, e)

        result = [
            f'✅ Created:          {created}',
            f'⚠️  Skipped (dup):   {skipped}',
            f'❌ Errors:           {errors}',
        ]
        if error_lines:
            result += ['', '--- Errors ---'] + error_lines[:20]

        self.result_text = '\n'.join(result)
        self.state = 'done'
        return self._reopen()

    def action_back(self):
        self.state = 'upload'
        return self._reopen()

    def action_close(self):
        return {'type': 'ir.actions.act_window_close'}

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
